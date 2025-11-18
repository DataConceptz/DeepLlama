"""
Export service for generating Markdown and DOCX reports
"""

import os
from typing import List, Dict
from datetime import datetime
from pathlib import Path


class ExportService:
    """Service for exporting reports in various formats"""

    def __init__(self):
        """Initialize export service"""
        pass

    def export_to_markdown(self, content: str, filepath: str, metadata: Dict = None) -> bool:
        """Export report to Markdown file"""
        try:
            # Add metadata header if provided
            if metadata:
                header = self._create_markdown_header(metadata)
                content = header + "\n\n" + content

            # Ensure directory exists
            Path(filepath).parent.mkdir(parents=True, exist_ok=True)

            # Write to file
            with open(filepath, 'w', encoding='utf-8') as f:
                f.write(content)

            return True

        except Exception as e:
            print(f"Error exporting to Markdown: {e}")
            return False

    def export_to_docx(self, content: str, filepath: str, metadata: Dict = None) -> bool:
        """Export report to DOCX file"""
        try:
            from docx import Document
            from docx.shared import Pt, Inches
            from docx.enum.text import WD_ALIGN_PARAGRAPH

            doc = Document()

            # Set default font
            style = doc.styles['Normal']
            font = style.font
            font.name = 'Times New Roman'
            font.size = Pt(12)

            # Add metadata/title page if provided
            if metadata:
                title = metadata.get('title', 'Literature Review Report')
                doc.add_heading(title, 0)

                if metadata.get('author'):
                    p = doc.add_paragraph(f"Author: {metadata['author']}")
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

                if metadata.get('date'):
                    p = doc.add_paragraph(f"Date: {metadata['date']}")
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

                doc.add_page_break()

            # Parse markdown-style content and add to document
            lines = content.split('\n')
            i = 0
            while i < len(lines):
                line = lines[i].strip()

                if not line:
                    i += 1
                    continue

                # Headers
                if line.startswith('# '):
                    doc.add_heading(line[2:], level=1)
                elif line.startswith('## '):
                    doc.add_heading(line[3:], level=2)
                elif line.startswith('### '):
                    doc.add_heading(line[4:], level=3)
                elif line.startswith('#### '):
                    doc.add_heading(line[5:], level=4)
                # Lists
                elif line.startswith('- ') or line.startswith('* '):
                    doc.add_paragraph(line[2:], style='List Bullet')
                elif line.startswith('1. ') or line.startswith('2. '):
                    # Numbered list
                    doc.add_paragraph(line[3:], style='List Number')
                # Code blocks
                elif line.startswith('```'):
                    i += 1
                    code_lines = []
                    while i < len(lines) and not lines[i].strip().startswith('```'):
                        code_lines.append(lines[i])
                        i += 1
                    code_text = '\n'.join(code_lines)
                    p = doc.add_paragraph(code_text)
                    p.style = 'No Spacing'
                    font = p.runs[0].font
                    font.name = 'Courier New'
                    font.size = Pt(10)
                # Regular paragraph
                else:
                    # Handle inline formatting
                    para = doc.add_paragraph()
                    self._add_formatted_text(para, line)

                i += 1

            # Ensure directory exists
            Path(filepath).parent.mkdir(parents=True, exist_ok=True)

            # Save document
            doc.save(filepath)
            return True

        except ImportError:
            print("python-docx not installed. Install with: pip install python-docx")
            return False
        except Exception as e:
            print(f"Error exporting to DOCX: {e}")
            return False

    def _create_markdown_header(self, metadata: Dict) -> str:
        """Create markdown metadata header"""
        header = "---\n"
        for key, value in metadata.items():
            header += f"{key}: {value}\n"
        header += "---"
        return header

    def _add_formatted_text(self, paragraph, text: str):
        """Add text with markdown-style formatting to paragraph"""
        # Simple implementation - can be enhanced
        # Handle bold **text** and italic *text*

        import re

        # Replace **bold**
        parts = re.split(r'(\*\*.*?\*\*)', text)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = paragraph.add_run(part[2:-2])
                run.bold = True
            elif part.startswith('*') and part.endswith('*') and not part.startswith('**'):
                run = paragraph.add_run(part[1:-1])
                run.italic = True
            else:
                paragraph.add_run(part)

    def create_report_template(self, title: str, query: str, articles: List[Dict]) -> str:
        """Create a basic report template"""
        report = f"""# {title}

## Executive Summary

*Query: {query}*

This literature review examines {len(articles)} articles related to {query}.

## Methodology

Articles were retrieved from multiple academic databases including OpenAlex, PubMed, arXiv, and CORE.

## Key Findings

[To be generated by AI model]

## Discussion

[To be generated by AI model]

## Conclusions

[To be generated by AI model]

## References

"""
        for i, article in enumerate(articles, 1):
            report += f"\n[{i}] {article.get('authors', 'Unknown')} ({article.get('year', 'n.d.')}). "
            report += f"{article.get('title', 'Untitled')}. "
            if article.get('journal'):
                report += f"*{article['journal']}*. "
            if article.get('doi'):
                report += f"https://doi.org/{article['doi']}"
            report += "\n"

        return report

    def save_articles_csv(self, articles: List[Dict], filepath: str) -> bool:
        """Export articles to CSV format"""
        try:
            import csv

            # Ensure directory exists
            Path(filepath).parent.mkdir(parents=True, exist_ok=True)

            with open(filepath, 'w', newline='', encoding='utf-8') as f:
                if not articles:
                    return False

                # Define fields
                fieldnames = ['title', 'authors', 'year', 'journal', 'doi', 'keywords',
                            'abstract', 'source', 'score', 'url', 'date_added', 'notes']

                writer = csv.DictWriter(f, fieldnames=fieldnames)
                writer.writeheader()

                for article in articles:
                    row = {k: article.get(k, '') for k in fieldnames}
                    writer.writerow(row)

            return True

        except Exception as e:
            print(f"Error exporting to CSV: {e}")
            return False

    def load_articles_csv(self, filepath: str) -> List[Dict]:
        """Import articles from CSV file"""
        try:
            import csv

            articles = []

            with open(filepath, 'r', encoding='utf-8') as f:
                reader = csv.DictReader(f)

                for row in reader:
                    # Convert year to int if possible
                    if row.get('year'):
                        try:
                            row['year'] = int(row['year'])
                        except:
                            row['year'] = 0

                    # Convert score to float if possible
                    if row.get('score'):
                        try:
                            row['score'] = float(row['score'])
                        except:
                            row['score'] = 0.0

                    articles.append(row)

            return articles

        except Exception as e:
            print(f"Error importing from CSV: {e}")
            return []
